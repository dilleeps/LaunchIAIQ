"""PRD .xlsx import — matches the layout of the Drug_Launch_PRD_Workbook.xlsx template.

Accepts a multipart upload. Reads the PRD_Template sheet (single-launch) OR the
Portfolio_Tracker sheet (one row per Asset×Country launch). Creates Asset, Country,
Launch, PRD, Milestones, Forecast rows. Idempotent on `launch_code`.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import get_current_user, require_role
from ..models import (
    Asset,
    Country,
    Forecast,
    Launch,
    Milestone,
    MilestoneTemplate,
    Organization,
    PRD,
    PRDVersion,
    User,
    Workstream,
)


router = APIRouter(prefix="/prd", tags=["prd-import"])


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _to_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, date):
        return v
    if isinstance(v, datetime):
        return v.date()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d %b %Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_float(v) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def _to_int(v) -> Optional[int]:
    f = _to_float(v)
    return int(f) if f is not None else None


def _get_or_create_asset(db: Session, org_id: uuid.UUID, brand_name: str, ta: str | None = None, molecule: str | None = None) -> Asset:
    a = db.query(Asset).filter(Asset.org_id == org_id, Asset.brand_name == brand_name).first()
    if a:
        return a
    a = Asset(org_id=org_id, brand_name=brand_name, therapeutic_area=ta, molecule_type=molecule)
    db.add(a)
    db.flush()
    return a


def _get_or_create_country(db: Session, name: str) -> Optional[Country]:
    if not name:
        return None
    c = db.query(Country).filter(Country.name == name).first()
    if c:
        return c
    # Auto-create with best-effort code
    code = (name[:3] or "ZZZ").upper()
    c = Country(code=code, name=name)
    db.add(c)
    db.flush()
    return c


def _parse_portfolio_tracker(ws, db: Session, org_id: uuid.UUID, user_id: uuid.UUID) -> list[Launch]:
    """One row per launch. Expected columns (header in row 3):
    Launch ID | Asset / Brand | Country | Region | Therapeutic Area | Indication |
    Launch Type | Launch Phase | Target Launch Date | Reg. Approval Date | HTA Decision |
    Overall RAG | Y1 Patients (Fcst) | Y1 Revenue (Local Curr.) | Country Launch Lead | ...
    """
    created: list[Launch] = []
    rows = list(ws.iter_rows(values_only=True))
    # Find header row
    header_idx = None
    for i, row in enumerate(rows[:10]):
        if row and any(_norm(c).lower() == "launch id" for c in row):
            header_idx = i
            break
    if header_idx is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Portfolio_Tracker: header row not found")
    header = [_norm(c).lower() for c in rows[header_idx]]
    def col(name: str) -> int:
        try:
            return header.index(name.lower())
        except ValueError:
            return -1

    for r in rows[header_idx + 1 :]:
        code = _norm(r[col("launch id")]) if col("launch id") >= 0 else ""
        if not code or code.startswith("[") or code.startswith("e.g"):
            continue
        # Skip if already imported
        if db.query(Launch).filter(Launch.org_id == org_id, Launch.launch_code == code).first():
            continue
        brand = _norm(r[col("asset / brand")]) if col("asset / brand") >= 0 else ""
        if not brand:
            continue
        country_name = _norm(r[col("country")]) if col("country") >= 0 else ""
        ta = _norm(r[col("therapeutic area")]) if col("therapeutic area") >= 0 else None
        asset = _get_or_create_asset(db, org_id, brand, ta=ta)
        country = _get_or_create_country(db, country_name)
        if not country:
            continue
        launch = Launch(
            org_id=org_id,
            launch_code=code,
            asset_id=asset.id,
            country_id=country.id,
            launch_type=_norm(r[col("launch type")]) or None,
            launch_phase=_norm(r[col("launch phase")]) or None,
            target_launch_date=_to_date(r[col("target launch date")]) if col("target launch date") >= 0 else None,
            reg_approval_date=_to_date(r[col("reg. approval date")]) if col("reg. approval date") >= 0 else None,
            hta_decision=_norm(r[col("hta decision")]) or None,
            overall_rag=_norm(r[col("overall rag")]) or "Green",
            country_launch_lead_user_id=user_id,
            global_brand_lead_user_id=user_id,
        )
        db.add(launch)
        db.flush()

        # Forecast (Y1 only from this sheet; Y2/Y3 unset)
        y1p = _to_int(r[col("y1 patients (fcst)")]) if col("y1 patients (fcst)") >= 0 else None
        y1r = _to_float(r[col("y1 revenue (local curr.)")]) if col("y1 revenue (local curr.)") >= 0 else None
        if y1p or y1r:
            y1_price = (y1r / y1p) if (y1p and y1r) else None
            db.add(
                Forecast(
                    launch_id=launch.id,
                    currency=country.currency_code or "USD",
                    source="bottom_up",
                    y1_patients=y1p,
                    y1_net_price=y1_price,
                )
            )

        # Apply default NCE milestone template if exists
        templates = db.query(MilestoneTemplate).filter(
            MilestoneTemplate.org_id == org_id,
            MilestoneTemplate.launch_type == (launch.launch_type or "NCE"),
        ).all()
        offset = 0
        for tpl in templates:
            tgt = launch.target_launch_date
            if tgt:
                from datetime import timedelta
                tgt = tgt + timedelta(days=tpl.default_offset_days or offset)
            db.add(
                Milestone(
                    launch_id=launch.id,
                    template_id=tpl.id,
                    workstream_id=tpl.workstream_id,
                    name=tpl.name,
                    target_date=tgt,
                    status="Not Started",
                    weight=tpl.weight,
                    is_gate=tpl.is_gate,
                )
            )
            offset += 14

        # Seed an empty PRD
        prd = PRD(launch_id=launch.id, current_version=1)
        db.add(prd)
        db.flush()
        db.add(
            PRDVersion(
                prd_id=prd.id,
                version=1,
                is_baseline=True,
                payload={
                    "identification": {
                        "launch_id": code,
                        "asset": brand,
                        "country": country.name,
                        "launch_type": launch.launch_type,
                        "launch_phase": launch.launch_phase,
                    },
                },
                created_by=user_id,
            )
        )
        created.append(launch)
    return created


def _parse_prd_template(ws, db: Session, org_id: uuid.UUID, user_id: uuid.UUID) -> Optional[Launch]:
    """The PRD_Template sheet is label/value pairs (col A label, col B value)."""
    kv: dict[str, str] = {}
    section = "identification"
    payload: dict[str, dict] = {section: {}}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        label = _norm(row[0])
        value = _norm(row[1]) if len(row) > 1 else ""
        if label.startswith(tuple("1234567890")) and "." in label:
            # New section header e.g. "2. Executive Summary"
            section = label.split(".", 1)[1].strip().lower().replace(" ", "_")
            payload.setdefault(section, {})
            continue
        if label and value and not value.startswith("["):
            key = label.lower().replace(" ", "_").replace("(", "").replace(")", "").replace("/", "_").replace("'", "")[:60]
            payload[section][key] = value
            kv[label] = value

    code = kv.get("Launch ID") or kv.get("launch_id") or f"L-{uuid.uuid4().hex[:6].upper()}"
    if db.query(Launch).filter(Launch.org_id == org_id, Launch.launch_code == code).first():
        raise HTTPException(status.HTTP_409_CONFLICT, f"Launch {code} already exists")
    brand = kv.get("Asset / Brand Name") or kv.get("Brand Name (Country)") or "Unnamed Asset"
    country_name = kv.get("Country / Market") or kv.get("Country") or "United States"
    ta = kv.get("Therapeutic Area")

    asset = _get_or_create_asset(db, org_id, brand, ta=ta, molecule=kv.get("Molecule Type"))
    country = _get_or_create_country(db, country_name)
    if not country:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Country missing")

    launch = Launch(
        org_id=org_id,
        launch_code=code,
        asset_id=asset.id,
        country_id=country.id,
        launch_type=kv.get("Launch Type"),
        launch_phase=kv.get("Launch Phase"),
        target_launch_date=_to_date(kv.get("Target Launch Date")),
        reg_approval_date=_to_date(kv.get("Approval Date (Actual/Expected)")),
        hta_decision=kv.get("HTA Decision"),
        country_launch_lead_user_id=user_id,
        global_brand_lead_user_id=user_id,
    )
    db.add(launch)
    db.flush()

    prd = PRD(launch_id=launch.id, current_version=1)
    db.add(prd)
    db.flush()
    db.add(
        PRDVersion(prd_id=prd.id, version=1, is_baseline=True, payload=payload, created_by=user_id)
    )
    return launch


@router.post("/import", dependencies=[Depends(require_role("global_admin", "global_brand_lead"))])
async def import_prd(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Upload an .xlsx file")
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, "openpyxl not installed on server")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Could not read xlsx: {exc}")

    created: list[Launch] = []
    if "Portfolio_Tracker" in wb.sheetnames:
        created = _parse_portfolio_tracker(wb["Portfolio_Tracker"], db, user.org_id, user.id)
    elif "PRD_Template" in wb.sheetnames:
        ln = _parse_prd_template(wb["PRD_Template"], db, user.org_id, user.id)
        if ln:
            created = [ln]
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Workbook must contain a Portfolio_Tracker or PRD_Template sheet",
        )

    db.commit()
    return {
        "imported": len(created),
        "launches": [
            {"id": str(ln.id), "launch_code": ln.launch_code, "brand": ln.asset.brand_name, "country": ln.country.name}
            for ln in created
        ],
    }
